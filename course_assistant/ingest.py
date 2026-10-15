"""Document parsing, page rendering and chunking.

Every chunk keeps its source details (document name, page/slide number, image path)
so answers can cite the exact page and show the original slide image.

Supported formats:
- .pdf  : text per page (PyMuPDF) + rendered page image
- .pptx : converted to PDF with LibreOffice (headless), then handled as PDF; slide text
          from python-pptx is used as a fallback for image-only slides
- .docx : converted to PDF with LibreOffice, then handled as PDF; falls back to text-only
          extraction with python-docx if LibreOffice is unavailable
- .txt / .md : plain text, split into pseudo-pages, no images
"""
from __future__ import annotations

import hashlib
import os
import shutil
import subprocess
import tempfile
from dataclasses import dataclass
from pathlib import Path

import pymupdf as fitz
from langchain_text_splitters import RecursiveCharacterTextSplitter

SUPPORTED_EXTENSIONS = (".pdf", ".pptx", ".docx", ".txt", ".md")
RENDER_DPI = 110
TEXT_PAGE_CHARS = 3000


class UnsupportedFormat(ValueError):
    pass


class ConversionError(RuntimeError):
    pass


@dataclass
class Page:
    doc_id: str
    doc_name: str
    page_number: int  # 1-based page / slide number
    text: str
    image_path: Path | None = None


@dataclass
class Chunk:
    chunk_id: str
    doc_id: str
    doc_name: str
    page_number: int
    text: str
    image_path: Path | None = None
    chunk_index: int = 0  # position within the page

    def to_dict(self) -> dict:
        return {
            "chunk_id": self.chunk_id,
            "doc_id": self.doc_id,
            "doc_name": self.doc_name,
            "page_number": self.page_number,
            "text": self.text,
            "image_path": str(self.image_path) if self.image_path else None,
            "chunk_index": self.chunk_index,
        }

    @classmethod
    def from_dict(cls, d: dict) -> "Chunk":
        return cls(
            chunk_id=d["chunk_id"],
            doc_id=d["doc_id"],
            doc_name=d["doc_name"],
            page_number=int(d["page_number"]),
            text=d["text"],
            image_path=Path(d["image_path"]) if d.get("image_path") else None,
            chunk_index=int(d.get("chunk_index", 0)),
        )


def file_sha256(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for block in iter(lambda: f.read(1 << 20), b""):
            h.update(block)
    return h.hexdigest()


def describe_supported_formats() -> str:
    return (
        "Accepted: PDF (.pdf), PowerPoint (.pptx), Word (.docx), text (.txt, .md). "
        "PPTX and DOCX are converted to PDF with LibreOffice so each slide/page is "
        "rendered as an image; text and original page images are both preserved."
    )


# ----------------------------------------------------------------------------
# Conversion
# ----------------------------------------------------------------------------

# Where LibreOffice lives when it is installed but not on PATH (macOS app bundle,
# Linux packages, Windows default). PATH is checked first.
_SOFFICE_CANDIDATES = (
    "/Applications/LibreOffice.app/Contents/MacOS/soffice",
    "/usr/lib/libreoffice/program/soffice",
    "/opt/libreoffice/program/soffice",
    r"C:\Program Files\LibreOffice\program\soffice.exe",
)


def find_libreoffice() -> str | None:
    """Return the soffice executable, honouring $SOFFICE_PATH, PATH, then known locations."""
    override = os.environ.get("SOFFICE_PATH")
    if override and Path(override).is_file():
        return override
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    if exe:
        return exe
    for cand in _SOFFICE_CANDIDATES:
        if Path(cand).is_file():
            return cand
    return None


def libreoffice_available() -> bool:
    return find_libreoffice() is not None


def convert_to_pdf(src: Path, out_dir: Path, timeout: int = 180) -> Path:
    """Convert an Office document to PDF with LibreOffice headless."""
    exe = find_libreoffice()
    if not exe:
        raise ConversionError("LibreOffice (soffice) is not installed; export the file to PDF manually.")
    out_dir.mkdir(parents=True, exist_ok=True)
    # Isolated profile dir avoids lock conflicts with a running LibreOffice.
    with tempfile.TemporaryDirectory(prefix="lo-profile-") as profile:
        cmd = [
            exe, "--headless", "--norestore", f"-env:UserInstallation=file://{profile}",
            "--convert-to", "pdf", "--outdir", str(out_dir), str(src),
        ]
        try:
            proc = subprocess.run(cmd, capture_output=True, text=True, timeout=timeout)
        except subprocess.TimeoutExpired:
            raise ConversionError(f"LibreOffice timed out converting {src.name}") from None
    pdf = out_dir / (src.stem + ".pdf")
    if proc.returncode != 0 or not pdf.exists():
        raise ConversionError(f"LibreOffice failed to convert {src.name}: {proc.stderr.strip()[:300]}")
    return pdf


# ----------------------------------------------------------------------------
# Parsing
# ----------------------------------------------------------------------------

def _clean(text: str) -> str:
    lines = [ln.rstrip() for ln in text.replace("\r", "").split("\n")]
    out: list[str] = []
    for ln in lines:
        if ln.strip() == "" and out and out[-1] == "":
            continue
        out.append(ln.strip())
    return "\n".join(out).strip()


def parse_pdf(pdf_path: Path, doc_id: str, doc_name: str, pages_dir: Path,
              fallback_text: dict[int, str] | None = None) -> list[Page]:
    pages_dir.mkdir(parents=True, exist_ok=True)
    pages: list[Page] = []
    zoom = RENDER_DPI / 72.0
    with fitz.open(pdf_path) as doc:
        for i, page in enumerate(doc, start=1):
            text = _clean(page.get_text("text"))
            if not text and fallback_text and fallback_text.get(i):
                text = _clean(fallback_text[i])
            img_path = pages_dir / f"page_{i:04d}.png"
            page.get_pixmap(matrix=fitz.Matrix(zoom, zoom), alpha=False).save(str(img_path))
            # Sidecar text next to the image: used by offline mock visual embeddings and
            # as a human-readable record of what the image contains.
            img_path.with_suffix(".txt").write_text(text, encoding="utf-8")
            pages.append(Page(doc_id, doc_name, i, text, img_path))
    return pages


def _pptx_slide_text(pptx_path: Path) -> dict[int, str]:
    from pptx import Presentation

    out: dict[int, str] = {}
    prs = Presentation(str(pptx_path))
    for i, slide in enumerate(prs.slides, start=1):
        parts = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text)
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append("Speaker notes: " + notes)
        out[i] = "\n".join(parts)
    return out


def parse_pptx(path: Path, doc_id: str, doc_name: str, pages_dir: Path) -> list[Page]:
    slide_text = _pptx_slide_text(path)
    with tempfile.TemporaryDirectory(prefix="pptx2pdf-") as tmp:
        pdf = convert_to_pdf(path, Path(tmp))
        pages = parse_pdf(pdf, doc_id, doc_name, pages_dir, fallback_text=slide_text)
    # Append speaker notes (not rendered in the PDF) to the slide text.
    for p in pages:
        notes = slide_text.get(p.page_number, "")
        if "Speaker notes:" in notes:
            note = notes.split("Speaker notes:", 1)[1].strip()
            if note and note not in p.text:
                p.text = (p.text + "\n\nSpeaker notes: " + note).strip()
                p.image_path.with_suffix(".txt").write_text(p.text, encoding="utf-8")
    return pages


def parse_docx(path: Path, doc_id: str, doc_name: str, pages_dir: Path) -> list[Page]:
    if libreoffice_available():
        try:
            with tempfile.TemporaryDirectory(prefix="docx2pdf-") as tmp:
                pdf = convert_to_pdf(path, Path(tmp))
                return parse_pdf(pdf, doc_id, doc_name, pages_dir)
        except ConversionError:
            pass  # fall through to text-only extraction
    import docx

    d = docx.Document(str(path))
    text = "\n".join(p.text for p in d.paragraphs)
    for table in d.tables:
        for row in table.rows:
            text += "\n" + " | ".join(c.text for c in row.cells)
    return _text_pages(_clean(text), doc_id, doc_name)


def _text_pages(text: str, doc_id: str, doc_name: str) -> list[Page]:
    if not text:
        return [Page(doc_id, doc_name, 1, "")]
    pages, start, n = [], 0, 1
    while start < len(text):
        end = min(len(text), start + TEXT_PAGE_CHARS)
        if end < len(text):
            nl = text.rfind("\n", start, end)
            if nl > start + TEXT_PAGE_CHARS // 2:
                end = nl
        pages.append(Page(doc_id, doc_name, n, text[start:end].strip()))
        start, n = end, n + 1
    return pages


def parse_document(path: Path, doc_id: str, doc_name: str, pages_dir: Path) -> list[Page]:
    ext = path.suffix.lower()
    if ext == ".pdf":
        return parse_pdf(path, doc_id, doc_name, pages_dir)
    if ext == ".pptx":
        return parse_pptx(path, doc_id, doc_name, pages_dir)
    if ext == ".docx":
        return parse_docx(path, doc_id, doc_name, pages_dir)
    if ext in (".txt", ".md"):
        return _text_pages(_clean(path.read_text(encoding="utf-8", errors="replace")), doc_id, doc_name)
    raise UnsupportedFormat(f"Unsupported file type '{ext}'. {describe_supported_formats()}")


# ----------------------------------------------------------------------------
# Chunking (page-bounded so every chunk keeps an exact page/slide reference)
# ----------------------------------------------------------------------------

def chunk_pages(pages: list[Page], chunk_size: int = 800, chunk_overlap: int = 120) -> list[Chunk]:
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=chunk_size, chunk_overlap=chunk_overlap,
        separators=["\n\n", "\n", ". ", " ", ""],
    )
    chunks: list[Chunk] = []
    for page in pages:
        pieces = splitter.split_text(page.text) if page.text else []
        if not pieces and page.image_path is not None:
            # Image-only page: keep a placeholder chunk so the visual index and
            # citations can still reach this slide.
            pieces = [f"[{page.doc_name}, page {page.page_number}: no extractable text; see image]"]
        for j, piece in enumerate(pieces):
            chunks.append(Chunk(
                chunk_id=f"{page.doc_id}:p{page.page_number}:c{j}",
                doc_id=page.doc_id, doc_name=page.doc_name, page_number=page.page_number,
                text=piece, image_path=page.image_path, chunk_index=j,
            ))
    return chunks
