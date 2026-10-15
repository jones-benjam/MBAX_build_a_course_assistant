"""Helpers that generate small course-like documents for tests."""
from pathlib import Path

import pymupdf as fitz


def make_pdf(path: Path, pages: list[str]) -> Path:
    doc = fitz.open()
    for text in pages:
        page = doc.new_page(width=792, height=612)  # landscape slide-like page
        # Textbox wraps long text so every word stays on the page (and is extractable).
        page.insert_textbox(fitz.Rect(36, 36, 756, 576), text, fontsize=11)
    doc.save(str(path))
    doc.close()
    return path


def make_pptx(path: Path, slides: list[tuple[str, str]], notes: dict[int, str] | None = None) -> Path:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    for i, (title, body) in enumerate(slides, start=1):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
        if notes and i in notes:
            slide.notes_slide.notes_text_frame.text = notes[i]
    # An image-only slide (no text) to exercise the placeholder-chunk path.
    blank = prs.slides.add_slide(prs.slide_layouts[6])
    shape = blank.shapes.add_shape(1, Inches(1), Inches(1), Inches(4), Inches(3))
    shape.fill.solid()
    prs.save(str(path))
    return path


def make_docx(path: Path, paragraphs: list[str]) -> Path:
    import docx

    d = docx.Document()
    for p in paragraphs:
        d.add_paragraph(p)
    d.save(str(path))
    return path
