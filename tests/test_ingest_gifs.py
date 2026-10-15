"""Tests for animated-GIF handling in PPTX ingestion (issue #21).

LibreOffice renders only the first frame of an animated GIF, which on some
slides is blank/black. `_patch_gif_frames` pastes the GIF's most informative
frame over the rendered image so the page image and visual index show content.

The frame-selection and patching helpers are tested without LibreOffice; the
full PPTX path test is skipped when LibreOffice is not installed.
"""
from io import BytesIO
from pathlib import Path

import numpy as np
import pytest
from PIL import Image, ImageStat

from course_assistant import ingest
from course_assistant.ingest import Page

needs_libreoffice = pytest.mark.skipif(not ingest.libreoffice_available(),
                                       reason="LibreOffice not installed")


def make_blank_first_gif(path: Path) -> Path:
    """Two-frame GIF: frame 1 is blank white, frame 2 has a dark pattern."""
    blank = Image.new("RGB", (120, 120), "white")
    detailed = Image.new("RGB", (120, 120), "white")
    for x in range(20, 100):
        for y in range(20, 100):
            if (x + y) % 3 == 0:
                detailed.putpixel((x, y), (20, 20, 200))
    blank.save(path, save_all=True, append_images=[detailed], duration=200,
               loop=0, disposal=2)
    return path


def test_emu_to_px_uses_render_dpi():
    # One inch is 914400 EMU; at RENDER_DPI that is RENDER_DPI pixels.
    assert ingest._emu_to_px(914400) == ingest.RENDER_DPI
    assert ingest._emu_to_px(0) == 0
    assert ingest._emu_to_px(457200) == ingest.RENDER_DPI // 2


def test_best_gif_frame_skips_blank_first_frame(tmp_path):
    gif = make_blank_first_gif(tmp_path / "anim.gif")
    frame = ingest._best_gif_frame(gif.read_bytes())
    stat = ImageStat.Stat(frame.convert("L"))
    assert stat.stddev[0] > 5  # meaningful content, not a blank frame


def test_gif_pictures_discovers_gif_on_slide(tmp_path):
    from pptx import Presentation

    gif = make_blank_first_gif(tmp_path / "anim.gif")
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank layout
    slide.shapes.add_picture(str(gif), ingest._emu_to_px(457200),
                             ingest._emu_to_px(457200))
    pptx_path = tmp_path / "deck.pptx"
    prs.save(pptx_path)

    pics = ingest._gif_pictures(pptx_path)
    assert 1 in pics and len(pics[1]) == 1
    assert pics[1][0]["blob"].startswith(b"GIF8")
    assert pics[1][0]["width"] > 0 and pics[1][0]["height"] > 0


def test_patch_gif_frames_fills_blank_area(tmp_path):
    # A rendered page that is blank white (as LibreOffice produces for a GIF
    # whose first frame is blank)...
    page_png = tmp_path / "page_0001.png"
    Image.new("RGB", (400, 300), "white").save(page_png)
    gif = make_blank_first_gif(tmp_path / "anim.gif")
    pages = [Page(doc_id="d", doc_name="deck.pptx", page_number=1,
                  text="", image_path=page_png)]
    gifs = {1: [{"left": 0, "top": 0,
                 "width": 13_333_333, "height": 7_500_000,
                 "blob": gif.read_bytes()}]}
    ingest._patch_gif_frames(pages, gifs)

    img = Image.open(page_png).convert("L")
    arr = np.asarray(img)
    # The patched region (where we pasted) must no longer be blank.
    region = arr[50:250, 50:350]
    assert region.std() > 5


@needs_libreoffice
def test_pptx_with_gif_renders_meaningful_slide_image(tmp_path):
    from pptx import Presentation

    gif = make_blank_first_gif(tmp_path / "anim.gif")
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_picture(str(gif), 914400, 914400, 3 * 914400, 2 * 914400)
    pptx_path = tmp_path / "deck.pptx"
    prs.save(pptx_path)

    pages = ingest.parse_pptx(pptx_path, "doc_gif", "deck.pptx", tmp_path / "pages")
    assert pages and pages[0].image_path is not None
    arr = np.asarray(Image.open(pages[0].image_path).convert("L"))
    # Slide image must contain visible content where the GIF was patched.
    assert arr.std() > 5
